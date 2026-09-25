"""Threads 実データ(インサイト画面のスクリーンショットから転記)の分析シートを生成する。

使い方:
    python sns-analysis/threads_2026-09.py
    -> sns-analysis/threads_2026-09.xlsx
数値は 2026/09/24 時点の Threads インサイト画面から手で転記したもの。
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).with_name("threads_2026-09.xlsx")
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, bold=True, size=14)
NOTE = Font(name=FONT, italic=True, color="808080")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
SRC = "出典: Threads インサイト画面(ユーザー提供のスクリーンショット、2026/09/24 時点)"

# 投稿名, カテゴリ, 投稿からの経過, 閲覧数, いいね, 返信, 再投稿, シェア
POSTS = [
    ("沖縄限定ガチャシリーズ", "ガチャ", "2週間", 4374, 42, 0, 0, 0),
    ("人気食堂シリーズ 糸満市座波 魚太郎", "グルメ", "5週間", 3262, 25, 0, 1, 0),
    ("沖縄弁当屋シリーズ", "グルメ", "1週間", 3236, 11, 0, 0, 0),
    ("沖縄限定お店シリーズ(GRINGO TACOS)", "お店", "4週間", 1496, 3, 1, 1, 0),
    ("沖縄限定お店 おしゃれな店内、2Fには…", "お店", "3週間", 1408, 22, 0, 0, 0),
    ("沖縄 伝統芸能シリーズ", "伝統芸能", "4週間", 1130, 29, 0, 1, 0),
    ("沖縄限定お店シリーズ", "お店", "5週間", 1124, 14, 0, 0, 0),
    ("沖縄限定路上販売シリーズ", "レア・珍スポット", "1週間", 942, 6, 0, 0, 0),
    ("レア自動販売シリーズ ドライブ中にお漏らし", "レア・珍スポット", "3週間", 886, 11, 1, 0, 0),
    ("レアなテーマパークシリーズ", "レア・珍スポット", "5週間", 788, 7, 3, 0, 0),
    ("沖縄 伝統芸能シリーズ", "伝統芸能", "4週間", 689, 6, 0, 0, 0),
    ("沖縄限定レアたまごシリーズ", "グルメ", "4日", 677, 1, 0, 0, 0),
    ("沖縄限定お店シリーズ(最新)", "お店", "2時間", 177, 0, 1, 0, 0),
    ("沖縄限定居酒屋シリーズ", "グルメ", "1日", 287, 0, 1, 0, 0),
]
CATEGORIES = ["ガチャ", "グルメ", "お店", "伝統芸能", "レア・珍スポット"]


def header(ws, row, labels):
    for j, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=label)
        c.font, c.fill, c.alignment, c.border = HEAD, HEAD_FILL, CENTER, BOX


def build_posts(wb):
    ws = wb.active
    ws.title = "投稿一覧"
    header(ws, 1, ["投稿", "カテゴリ", "投稿からの経過", "閲覧数", "いいね", "返信", "再投稿", "シェア",
                   "反応合計", "いいね率", "反応率"])
    for i, p in enumerate(POSTS, start=2):
        for j, v in enumerate(p, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = BLUE, BOX
        ws[f"I{i}"] = f"=SUM(E{i}:H{i})"
        ws[f"J{i}"] = f"=IF(D{i}=0,0,E{i}/D{i})"
        ws[f"K{i}"] = f"=IF(D{i}=0,0,I{i}/D{i})"
        for col in "IJK":
            ws[f"{col}{i}"].font, ws[f"{col}{i}"].border = BLACK, BOX
        for col in "DEFGHI":
            ws[f"{col}{i}"].number_format = "#,##0"
        ws[f"J{i}"].number_format = ws[f"K{i}"].number_format = "0.00%"
    last = len(POSTS) + 1
    t = last + 1
    ws.cell(row=t, column=1, value="合計").font = BOLD
    for col in "DEFGHI":
        ws[f"{col}{t}"] = f"=SUM({col}2:{col}{last})"
        ws[f"{col}{t}"].number_format, ws[f"{col}{t}"].font = "#,##0", BOLD
    ws[f"J{t}"] = f"=IF(D{t}=0,0,E{t}/D{t})"
    ws[f"K{t}"] = f"=IF(D{t}=0,0,I{t}/D{t})"
    for col in "JK":
        ws[f"{col}{t}"].number_format, ws[f"{col}{t}"].font = "0.00%", BOLD
    ws.cell(row=t + 2, column=1, value=SRC).font = NOTE
    ws.cell(row=t + 3, column=1, value="※ 閲覧数上位12件(90日間)+最新投稿2件。カテゴリは投稿タイトルから分類。青字=転記した値、黒字=計算式。").font = NOTE
    for col, w in zip("ABCDEFGHIJK", [40, 16, 14, 10, 8, 8, 8, 8, 10, 10, 10]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    return last


def build_category(wb, last):
    ws = wb.create_sheet("カテゴリ別")
    ws["A1"] = "カテゴリ別の成績"
    ws["A1"].font = TITLE
    header(ws, 3, ["カテゴリ", "投稿数", "閲覧数合計", "1投稿あたり閲覧数", "いいね合計", "いいね率", "返信合計"])
    rng = lambda c: f"投稿一覧!${c}$2:${c}${last}"
    for i, cat in enumerate(CATEGORIES, start=4):
        ws[f"A{i}"] = cat
        ws[f"B{i}"] = f"=COUNTIFS({rng('B')},$A{i})"
        ws[f"C{i}"] = f"=SUMIFS({rng('D')},{rng('B')},$A{i})"
        ws[f"D{i}"] = f"=IF(B{i}=0,0,C{i}/B{i})"
        ws[f"E{i}"] = f"=SUMIFS({rng('E')},{rng('B')},$A{i})"
        ws[f"F{i}"] = f"=IF(C{i}=0,0,E{i}/C{i})"
        ws[f"G{i}"] = f"=SUMIFS({rng('F')},{rng('B')},$A{i})"
        for col in "ABCDEFG":
            c = ws[f"{col}{i}"]
            c.border = BOX
            c.font = BLACK if col in "ADF" else GREEN
            c.number_format = "0.00%" if col == "F" else "#,##0"
    end = 3 + len(CATEGORIES)
    for col, w in zip("ABCDEFG", [18, 8, 12, 18, 11, 10, 10]):
        ws.column_dimensions[col].width = w

    for col, title, fmt, anchor in ((4, "1投稿あたり閲覧数", "#,##0", "I3"), (6, "いいね率", "0.0%", "I19")):
        ch = BarChart()
        ch.type = "bar"
        ch.title = f"カテゴリ別 {title}"
        ch.legend = None
        ch.y_axis.number_format = fmt
        ch.height, ch.width = 7, 14
        ch.add_data(Reference(ws, min_col=col, min_row=3, max_row=end), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=end))
        ws.add_chart(ch, anchor)


def build_audience(wb):
    ws = wb.create_sheet("オーディエンス")
    ws["A1"] = "オーディエンス(直近30日: 8/26〜9/24)"
    ws["A1"].font = TITLE
    blocks = [
        ("概要", ["指標", "値"], [("閲覧数", 28830, "#,##0"), ("閲覧者", 20520, "#,##0"),
                                ("閲覧者のうちフォロワー", 12, "#,##0"), ("純フォロワー増", 12, "+#,##0"),
                                ("閲覧数の前期比", 3.975, "+0.0%"), ("閲覧者の前期比", 3.563, "+0.0%")]),
        ("年齢", ["年齢", "割合"], [("13-17", 0.0011, "0.0%"), ("18-24", 0.0626, "0.0%"), ("25-34", 0.304, "0.0%"),
                                  ("35-44", 0.3361, "0.0%"), ("45-54", 0.2114, "0.0%"), ("55-64", 0.0723, "0.0%"),
                                  ("65+", 0.0126, "0.0%")]),
        ("性別", ["性別", "割合"], [("女性", 0.3378, "0.0%"), ("男性", 0.1597, "0.0%"), ("その他・不明", 0.5025, "0.0%")]),
        ("市区町村(上位5)", ["市区町村", "割合"], [("那覇市", 0.0986, "0.0%"), ("沖縄市", 0.0765, "0.0%"),
                                              ("宜野湾市", 0.0472, "0.0%"), ("島尻郡", 0.0328, "0.0%"),
                                              ("浦添市", 0.0321, "0.0%")]),
        ("国・地域(上位5)", ["国・地域", "割合"], [("日本", 0.7773, "0.0%"), ("台湾", 0.0033, "0.0%"),
                                              ("香港", 0.0013, "0.0%"), ("アメリカ合衆国", 0.0012, "0.0%"),
                                              ("韓国", 0.0006, "0.0%")]),
        ("興味・関心", ["コミュニティ", "人数"], [("沖縄", 215, "#,##0"), ("Cats of Threads", 40, "#,##0"),
                                           ("旅行", 35, "#,##0"), ("筋トレ", 33, "#,##0"), ("沖縄グルメ", 28, "#,##0")]),
        ("最もアクティブな時間(JST)", ["曜日・時間", "閲覧者"], [("水曜 21時〜24時", 7715, "#,##0"),
                                                         ("火曜 21時〜24時", 6988, "#,##0"),
                                                         ("月曜 21時〜24時", 6573, "#,##0")]),
    ]
    r = 3
    for title, labels, rows in blocks:
        ws.cell(row=r, column=1, value=title).font = BOLD
        for j, label in enumerate(labels, start=1):
            c = ws.cell(row=r + 1, column=j, value=label)
            c.font, c.fill, c.alignment, c.border = HEAD, HEAD_FILL, CENTER, BOX
        for i, (k, v, fmt) in enumerate(rows, start=r + 2):
            ws.cell(row=i, column=1, value=k).font = BLACK
            c = ws.cell(row=i, column=2, value=v)
            c.font, c.number_format = BLUE, fmt
            ws.cell(row=i, column=1).border = c.border = BOX
        r += len(rows) + 3
    ws.cell(row=r, column=1, value=SRC).font = NOTE
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 12


def main():
    wb = Workbook()
    last = build_posts(wb)
    build_category(wb, last)
    build_audience(wb)
    wb.save(OUT)
    print(f"saved: {OUT}")


if __name__ == "__main__":
    main()
